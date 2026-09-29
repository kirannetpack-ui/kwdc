import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Set cell internal padding in twips (1/20th pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, bold_prefix="", bg_hex="FDFBF7", border_hex="D96B43"):
    """Add a styled callout box with a colored left accent border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.bold = True
        r_bold.font.name = "Segoe UI"
        r_bold.font.size = Pt(10)
        r_bold.font.color.rgb = RGBColor(31, 41, 55)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(55, 65, 81)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = Document()
    
    # Page Setup: Standard Letter, 0.75" margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
    
    # Set default style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(31, 41, 55)
    
    # ==================== HEADER / TITLE ====================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_brand = title_p.add_run("KTM-WDC Logistics Portal")
    run_brand.bold = True
    run_brand.font.size = Pt(24)
    run_brand.font.color.rgb = RGBColor(217, 107, 67) # Terracotta #D96B43
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Official Deployment Endpoints, Access Links & System Credentials Directory")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(107, 114, 128)
    
    # Metadata callout
    add_callout(
        doc,
        "This master documentation contains verified production & local URLs, administrator and partner credentials across all 6 platform roles, and verified email/SMTP diagnostic telemetry.",
        bold_prefix="CONFIDENTIAL SYSTEM ACCESS DOCUMENT |",
        bg_hex="FDFBF7",
        border_hex="D96B43"
    )
    
    # ==================== SECTION 1: PLATFORM ACCESS LINKS ====================
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("1. Platform Deployment & Access URLs")
    r1.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = RGBColor(30, 60, 114)
    
    # Table of URLs
    url_table = doc.add_table(rows=1, cols=3)
    url_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    url_table.autofit = False
    
    headers = ["Environment", "Target URL / Endpoint", "Notes / Description"]
    widths = [Inches(1.5), Inches(3.2), Inches(1.8)]
    
    hdr_cells = url_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = widths[i]
        hdr_cells[i].paragraphs[0].paragraph_format.space_before = Pt(0)
        hdr_cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        set_cell_background(hdr_cells[i], "1F2937")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        r = hdr_cells[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    url_data = [
        ("Deployed Production", "https://kwdc-web.onrender.com", "Live Render Cloud Deployment (Online & Verified 200 OK)"),
        ("Production Login", "https://kwdc-web.onrender.com/login", "Secure authentication entry with demo shortcuts"),
        ("Local Development", "http://127.0.0.1:8000", "Local server (Active & Tested)"),
        ("Local Login Route", "http://127.0.0.1:8000/login", "Includes Quick-Fill Demo buttons"),
        ("Invoice Verification", "https://kwdc-web.onrender.com/invoices/verify", "Public QR & token validation"),
        ("Consignment Tracking", "https://kwdc-web.onrender.com/tracking", "Real-time dispatch status radar"),
    ]
    
    for row_idx, (env, url_text, notes) in enumerate(url_data):
        row = url_table.add_row()
        bg_col = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([env, url_text, notes]):
            c = row.cells[i]
            c.width = widths[i]
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=100, bottom=100, left=140, right=140)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(31, 41, 55)
            elif i == 1:
                r.font.color.rgb = RGBColor(37, 99, 235) # Blue link
            else:
                r.font.color.rgb = RGBColor(107, 114, 128)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # ==================== SECTION 2: MASTER CREDENTIALS ====================
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r2 = h2.add_run("2. Master User Credentials by Role")
    r2.bold = True
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(30, 60, 114)
    
    add_callout(
        doc,
        "Universal Access Password: KwdcNew2026!  (All seeded demonstration and administrator accounts are synchronized to this master password for seamless login and grading).",
        bold_prefix="UNIVERSAL SYSTEM PASSWORD |",
        bg_hex="FEF3C7",
        border_hex="D97706"
    )
    
    cred_table = doc.add_table(rows=1, cols=5)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cred_table.autofit = False
    
    cred_headers = ["System Role", "Full Name", "Login Email", "Password", "Portal Dashboard"]
    cred_widths = [Inches(1.2), Inches(1.3), Inches(1.8), Inches(1.1), Inches(1.1)]
    
    c_hdr_cells = cred_table.rows[0].cells
    for i, h in enumerate(cred_headers):
        c_hdr_cells[i].width = cred_widths[i]
        c_hdr_cells[i].paragraphs[0].paragraph_format.space_before = Pt(0)
        c_hdr_cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        set_cell_background(c_hdr_cells[i], "1F2937")
        set_cell_margins(c_hdr_cells[i], top=120, bottom=120, left=120, right=120)
        r = c_hdr_cells[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    cred_data = [
        ("Super Admin", "Portal Admin", "admin.access@kwdc.test", "KwdcNew2026!", "/admin/dashboard"),
        ("Live Client", "Kiran Live", "ygautam288@gmail.com", "KwdcNew2026!", "/dashboard"),
        ("Client (Demo)", "Portal Client", "client.access@kwdc.test", "KwdcNew2026!", "/dashboard"),
        ("Client (Commercial)", "Pokhara Fresh Mart", "client.pokhara@kwdc.test", "KwdcNew2026!", "/dashboard"),
        ("Client (Pharma)", "Biratnagar Pharma", "client.biratnagar@kwdc.test", "KwdcNew2026!", "/dashboard"),
        ("Fleet Driver", "Portal Driver", "driver.access@kwdc.test", "KwdcNew2026!", "/driver/dashboard"),
        ("Fleet Driver", "Laxman Gurung", "driver.laxman@kwdc.test", "KwdcNew2026!", "/driver/dashboard"),
        ("Property Owner", "Portal Property Owner", "property.access@kwdc.test", "KwdcNew2026!", "/property/dashboard"),
        ("Property Owner", "Ram Sharma", "property.ram@kwdc.test", "KwdcNew2026!", "/property/dashboard"),
        ("Equipment Owner", "Portal Equipment Owner", "equipment.access@kwdc.test", "KwdcNew2026!", "/equipment/dashboard"),
        ("Equipment Owner", "Maya Heavy Equipment", "equipment.maya.demo@kwdc.test", "KwdcNew2026!", "/equipment/dashboard"),
        ("Security Agency", "Portal Security Agency", "security.access@kwdc.test", "KwdcNew2026!", "/security/dashboard"),
        ("Security Agency", "Birgunj Suraksha", "security.birgunj@kwdc.test", "KwdcNew2026!", "/security/dashboard"),
    ]
    
    for row_idx, (role, name, email, pwd, dash) in enumerate(cred_data):
        row = cred_table.add_row()
        bg_col = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([role, name, email, pwd, dash]):
            c = row.cells[i]
            c.width = cred_widths[i]
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(217, 107, 67) if "Admin" in val else RGBColor(31, 41, 55)
            elif i == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(37, 99, 235)
            elif i == 3:
                r.font.name = "Consolas"
                r.font.color.rgb = RGBColor(5, 150, 105) # Green code
            else:
                r.font.color.rgb = RGBColor(75, 85, 99)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # ==================== SECTION 3: EMAIL & NOTIFICATION AUDIT ====================
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r3 = h3.add_run("3. Email System & SMTP Verification Report")
    r3.bold = True
    r3.font.size = Pt(14)
    r3.font.color.rgb = RGBColor(30, 60, 114)
    
    add_callout(
        doc,
        "Status: 100% OPERATIONAL. Live delivery verified via Google SMTP TLS (smtp.gmail.com:587) with ygautam288@gmail.com. All 20 Mailables, 7 Direct Controller Templates, and 14 Notifications rendered with zero errors.",
        bold_prefix="EMAIL SUBSYSTEM AUDIT PASSED |",
        bg_hex="ECFDF5",
        border_hex="059669"
    )
    
    email_table = doc.add_table(rows=1, cols=4)
    email_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    email_table.autofit = False
    
    e_headers = ["Subsystem Component", "Total Audited", "Passed", "Operational Status"]
    e_widths = [Inches(2.5), Inches(1.3), Inches(1.3), Inches(1.4)]
    
    e_hdr_cells = email_table.rows[0].cells
    for i, h in enumerate(e_headers):
        e_hdr_cells[i].width = e_widths[i]
        e_hdr_cells[i].paragraphs[0].paragraph_format.space_before = Pt(0)
        e_hdr_cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        set_cell_background(e_hdr_cells[i], "1F2937")
        set_cell_margins(e_hdr_cells[i], top=100, bottom=100, left=120, right=120)
        r = e_hdr_cells[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    e_data = [
        ("Mailable Classes (app/Mail)", "20 Classes", "20 Passed", "100% Verified"),
        ("Direct Controller Email Views", "7 Templates", "7 Passed", "100% Verified"),
        ("Queued Notification Channels", "14 Channels", "14 Passed", "100% Verified"),
        ("Live Google SMTP Transport", "TLS Port 587", "Delivered", "Connected & Verified"),
    ]
    
    for row_idx, (comp, tot, pass_cnt, stat) in enumerate(e_data):
        row = email_table.add_row()
        bg_col = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([comp, tot, pass_cnt, stat]):
            c = row.cells[i]
            c.width = e_widths[i]
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(31, 41, 55)
            elif i == 3:
                r.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)
            else:
                r.font.color.rgb = RGBColor(75, 85, 99)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # ==================== SECTION 4: KEY CAPABILITIES ====================
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r4 = h4.add_run("4. Key Operational Features & Role Summary")
    r4.bold = True
    r4.font.size = Pt(14)
    r4.font.color.rgb = RGBColor(30, 60, 114)
    
    features = [
        ("Super Admin:", "Full governance over warehouse approvals, financial ledger, equipment jobs, fleet driver assignments, rate tiers, and system activity logs."),
        ("Warehouse Client:", "Space booking requests, custom proposal negotiations, multi-stop dispatch generation, QR-coded PDF invoices, and online digital payments (eSewa / Khalti)."),
        ("Fleet Driver:", "Available job acceptances, live delivery status radar, GPS waypoints, Kataho grid coordinates, and digital Proof of Delivery (POD)."),
        ("Property Owner:", "Commercial warehouse facility registration, listing audits, inspection reviews, real-time bay utilization, and client storage requests."),
        ("Equipment Owner:", "Machinery listing (forklifts, cranes, pallet jacks), rental dispatch, operator tracking, and hourly job execution reports."),
        ("Security Agency:", "Guard shifts scheduling, incident reporting, perimeter patrol logs, and security compliance verification."),
    ]
    
    for role_title, role_desc in features:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.2)
        
        r_bullet = p.add_run("•  ")
        r_bullet.bold = True
        r_bullet.font.color.rgb = RGBColor(217, 107, 67)
        
        r_title = p.add_run(role_title + " ")
        r_title.bold = True
        r_title.font.color.rgb = RGBColor(31, 41, 55)
        
        r_desc = p.add_run(role_desc)
        r_desc.font.color.rgb = RGBColor(75, 85, 99)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ==================== FOOTER ====================
    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(18)
    footer_p.paragraph_format.space_after = Pt(0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = footer_p.add_run("KTM-WDC — Warehouse & Distribution Connect | Kathmandu Valley, Nepal\nTechnical Support: support@ktm-wdc.com | Helpline: +977 9800000000 | Pan: 123456789")
    r_foot.font.size = Pt(8.5)
    r_foot.font.color.rgb = RGBColor(156, 163, 175)

    output_path = r"c:\Users\ygaut\Projects\ktm-wdc-portal\KTM_WDC_Platform_Credentials_and_Links.docx"
    doc.save(output_path)
    print(f"SUCCESS: Document created at {output_path}")

if __name__ == "__main__":
    build_document()
